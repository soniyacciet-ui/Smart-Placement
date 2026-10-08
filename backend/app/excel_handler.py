"""
Excel handler for DRIVE-X student data import.
Validates .xlsx files and returns structured preview data.
"""
import io
import uuid
from typing import Dict, List, Tuple
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from . import models


# In-memory store for preview sessions (temp storage between preview → confirm)
# In production, use Redis or a DB table. For DRIVE-X, in-memory is fine.
PREVIEW_CACHE: Dict[str, dict] = {}


REQUIRED_COLUMNS = ["name", "register_number", "email", "cgpa", "department"]


def _normalize_header(h: str) -> str:
    """Convert 'Register Number' → 'register_number'"""
    return str(h).strip().lower().replace(" ", "_")


def parse_and_validate(file_bytes: bytes, department: str, db: Session) -> dict:
    """
    Reads Excel from bytes, validates rows, returns preview data.
    Does NOT save anything to the database.
    """
    try:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
    except Exception as e:
        raise ValueError(f"Could not read Excel file: {e}")

    # ---- Read header row ----
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel file is empty")

    header = [_normalize_header(c) for c in rows[0]]

    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    col_index = {name: header.index(name) for name in REQUIRED_COLUMNS}

    # ---- Fetch existing students in this department ----
    existing = {
        s.register_number: s
        for s in db.query(models.Student).filter(models.Student.department == department).all()
        if s.register_number
    }

    # ---- Also fetch ALL students globally (to prevent cross-dept duplicate reg numbers) ----
    global_existing = {
        s.register_number: s
        for s in db.query(models.Student).all()
        if s.register_number
    }

    preview_rows: List[dict] = []
    seen_in_file: set = set()

    total = 0
    valid = 0
    invalid = 0
    duplicates = 0
    new_students = 0
    updated_students = 0

    for row_idx, raw_row in enumerate(rows[1:], start=2):  # Start at row 2 (Excel row numbers)
        # Skip fully empty rows
        if not raw_row or all(c is None or str(c).strip() == "" for c in raw_row):
            continue

        total += 1

        def get(col_name):
            idx = col_index[col_name]
            return raw_row[idx] if idx < len(raw_row) else None

        name = str(get("name") or "").strip()
        reg_no = str(get("register_number") or "").strip()
        email = str(get("email") or "").strip()
        cgpa_raw = get("cgpa")
        dept = str(get("department") or "").strip().upper() or department.upper()

        # ---- Validate row ----
        error = None
        if not name:
            error = "Missing Name"
        elif not reg_no:
            error = "Missing Register Number"
        elif not email or "@" not in email:
            error = "Invalid Email"
        elif cgpa_raw is None or str(cgpa_raw).strip() == "":
            error = "Missing CGPA"
        else:
            try:
                cgpa = float(cgpa_raw)
                if cgpa < 0 or cgpa > 10:
                    error = "CGPA must be between 0 and 10"
            except (ValueError, TypeError):
                error = "Invalid CGPA"

        # ---- Department check ----
        if not error and dept.upper() != department.upper():
            error = f"Department mismatch (row: {dept}, expected: {department})"

        # ---- Duplicate within file ----
        if not error and reg_no in seen_in_file:
            error = "Duplicate Register Number in Excel"
            duplicates += 1

        # ---- Row status ----
        if error:
            invalid += 1
            preview_rows.append({
                "row_number": row_idx,
                "name": name or None,
                "register_number": reg_no or None,
                "email": email or None,
                "cgpa": float(cgpa_raw) if isinstance(cgpa_raw, (int, float)) else None,
                "department": dept,
                "status": "Invalid",
                "reason": error,
            })
            continue

        # ---- Check existing in DB ----
        seen_in_file.add(reg_no)
        if reg_no in global_existing:
            # Registered elsewhere (cross-department) → reject
            if global_existing[reg_no].department != department:
                invalid += 1
                preview_rows.append({
                    "row_number": row_idx,
                    "name": name,
                    "register_number": reg_no,
                    "email": email,
                    "cgpa": float(cgpa_raw),
                    "department": dept,
                    "status": "Invalid",
                    "reason": f"Register Number belongs to {global_existing[reg_no].department}",
                })
                continue

            # Registered in this dept → UPDATE
            updated_students += 1
            preview_rows.append({
                "row_number": row_idx,
                "name": name,
                "register_number": reg_no,
                "email": email,
                "cgpa": float(cgpa_raw),
                "department": dept,
                "status": "Update",
                "reason": "Existing student will be updated",
            })
        else:
            # Brand new student
            new_students += 1
            valid += 1
            preview_rows.append({
                "row_number": row_idx,
                "name": name,
                "register_number": reg_no,
                "email": email,
                "cgpa": float(cgpa_raw),
                "department": dept,
                "status": "Valid",
                "reason": None,
            })

    # ---- Cache the preview for confirm step ----
    token = str(uuid.uuid4())
    PREVIEW_CACHE[token] = {
        "department": department,
        "rows": preview_rows,
        "file_name": None,
    }

    return {
        "total_rows": total,
        "valid_rows": valid,
        "invalid_rows": invalid,
        "duplicate_rows": duplicates,
        "new_students": new_students,
        "updated_students": updated_students,
        "rows": preview_rows,
        "token": token,
    }


def commit_import(token: str, db: Session) -> dict:
    """
    Reads the cached preview and commits all Valid + Update rows to the database.
    Runs inside a single transaction.
    """
    cached = PREVIEW_CACHE.get(token)
    if not cached:
        raise ValueError("Preview session expired or invalid")

    department = cached["department"]
    rows = cached["rows"]

    inserted = 0
    updated = 0
    failed = 0
    skipped = 0
    errors = []

    try:
        for row in rows:
            status = row["status"]

            if status == "Invalid":
                failed += 1
                errors.append(row)
                continue

            if status == "Duplicate":
                skipped += 1
                continue

            reg_no = row["register_number"]

            if status == "Valid":
                # New student
                new_student = models.Student(
                    name=row["name"],
                    register_number=reg_no,
                    email=row["email"],
                    cgpa=row["cgpa"],
                    department=department,
                    skills="",  # to be filled later
                )
                db.add(new_student)
                inserted += 1

            elif status == "Update":
                # Update existing
                student = db.query(models.Student).filter(
                    models.Student.register_number == reg_no
                ).first()
                if student:
                    student.name = row["name"]
                    student.email = row["email"]
                    student.cgpa = row["cgpa"]
                    student.department = department
                    updated += 1
                else:
                    # Row was marked Update but student no longer exists → treat as insert
                    new_student = models.Student(
                        name=row["name"],
                        register_number=reg_no,
                        email=row["email"],
                        cgpa=row["cgpa"],
                        department=department,
                        skills="",
                    )
                    db.add(new_student)
                    inserted += 1

        db.commit()
        PREVIEW_CACHE.pop(token, None)
    except Exception as e:
        db.rollback()
        raise ValueError(f"Database error during import: {e}")

    return {
        "total": len(rows),
        "new_students": inserted,
        "updated_students": updated,
        "failed": failed,
        "skipped": skipped,
        "errors": errors,
    }